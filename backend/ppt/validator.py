import logging
from pptx import Presentation
from pathlib import Path

logger = logging.getLogger(__name__)

class PPTValidator:
    def __init__(self, expected_slide_count=15):
        self.expected_slide_count = expected_slide_count
        
    def validate(self, ppt_path: Path):
        logger.info(f"Validating {ppt_path}...")
        try:
            prs = Presentation(ppt_path)
        except Exception as e:
            raise ValueError(f"Failed to open generated PPT: {e}")
            
        if len(prs.slides) != self.expected_slide_count:
            raise ValueError(f"Validation Failed: Expected {self.expected_slide_count} slides, found {len(prs.slides)}.")
            
        logger.info("Validation Passed: Structural integrity confirmed.")
        return True
