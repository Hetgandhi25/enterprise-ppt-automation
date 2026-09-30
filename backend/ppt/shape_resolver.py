from pptx.enum.shapes import MSO_SHAPE_TYPE

def find_shape_by_id(slide, target_id_str: str):
    """
    Recursively searches a slide to find a shape matching the given string ID.
    """
    def _search(shapes):
        for shape in shapes:
            if str(shape.shape_id) == target_id_str:
                return shape
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                result = _search(shape.shapes)
                if result:
                    return result
        return None
        
    return _search(slide.shapes)
