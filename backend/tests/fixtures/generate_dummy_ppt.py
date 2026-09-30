from pptx import Presentation
from pptx.util import Inches
from pathlib import Path

def add_alt_text(shape, alt_text):
    if hasattr(shape, "_element"):
        try:
            cNvPr = shape._element.xpath('.//p:cNvPr')
            if cNvPr:
                cNvPr[0].set("descr", alt_text)
                cNvPr[0].set("title", alt_text)
            shape.name = alt_text
        except:
            shape.name = alt_text

def create_dummy_template():
    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    
    # ---------------- SLIDE 5: Inventory Snapshot ----------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    txBox = slide5.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tf = txBox.text_frame
    tf.text = "Inventory Snapshot for {{customer_name}}"
    
    rows, cols = 3, 3
    table_shape = slide5.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(8), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Customer"
    table.cell(0, 1).text = "Service"
    table.cell(0, 2).text = "Links"
    
    table.cell(1, 0).text = "{{customer}}"
    table.cell(1, 1).text = "{{service}}"
    table.cell(1, 2).text = "{{links}}"
    
    add_alt_text(table_shape, "{{inventory_table}}")

    # ---------------- SLIDE 6: Service Performance ----------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    txBox = slide6.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tf = txBox.text_frame
    tf.text = "Service Performance Summary - {{report_month}}"
    
    shape1 = slide6.shapes.add_shape(1, Inches(1), Inches(2), Inches(4), Inches(2))
    add_alt_text(shape1, "{{sla_trend_chart}}")
    
    shape2 = slide6.shapes.add_shape(1, Inches(5.5), Inches(2), Inches(4), Inches(2))
    add_alt_text(shape2, "{{ticket_trend_chart}}")

    # ---------------- SLIDE 7: Location Wise SLA ----------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    txBox = slide7.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    txBox.text_frame.text = "Location Wise SLA"
    
    t1_shape = slide7.shapes.add_table(2, 3, Inches(1), Inches(2), Inches(4), Inches(1.5))
    t1 = t1_shape.table
    t1.cell(0, 0).text = "Location"
    t1.cell(0, 1).text = "Link ID"
    t1.cell(0, 2).text = "SLA %"
    add_alt_text(t1_shape, "{{links_below_sla_table}}")
    
    t2_shape = slide7.shapes.add_table(2, 3, Inches(5.5), Inches(2), Inches(4), Inches(1.5))
    t2 = t2_shape.table
    t2.cell(0, 0).text = "Location"
    t2.cell(0, 1).text = "Duration"
    t2.cell(0, 2).text = "Reason"
    add_alt_text(t2_shape, "{{major_downtime_table}}")

    # ---------------- SLIDE 9: Incident Summary ----------------
    slide9 = prs.slides.add_slide(blank_slide_layout)
    txBox = slide9.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    txBox.text_frame.text = "Incident Summary"
    
    t3_shape = slide9.shapes.add_table(2, 4, Inches(1), Inches(2), Inches(8), Inches(2))
    t3 = t3_shape.table
    t3.cell(0, 0).text = "Location"
    t3.cell(0, 1).text = "Cable"
    t3.cell(0, 2).text = "Backhaul"
    t3.cell(0, 3).text = "Electricity"
    add_alt_text(t3_shape, "{{incident_summary_table}}")

    out_dir = Path(__file__).parent.parent.parent / "templates"
    out_dir.mkdir(parents=True, exist_ok=True)
    prs.save(out_dir / "ServiceReview.pptx")
    print("Dummy template created.")

if __name__ == "__main__":
    create_dummy_template()
