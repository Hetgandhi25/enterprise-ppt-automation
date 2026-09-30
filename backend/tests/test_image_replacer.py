from pptx import Presentation
from pptx.util import Inches
from backend.ppt.image_replacer import replace_image
import pytest

def test_image_replacer(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
    
    # create a dummy image
    img_path = tmp_path / "dummy.png"
    from PIL import Image
    img = Image.new('RGB', (100, 100), color = 'red')
    img.save(img_path)
    
    replace_image(slide, shape, img_path)
    
    # Original shape should be detached
    # New picture shape should be added
    assert len(slide.shapes) == 1
    assert slide.shapes[0].shape_type == 13 # msoPICTURE
