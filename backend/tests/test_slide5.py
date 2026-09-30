import pytest
from pptx import Presentation
from pptx.util import Inches
from backend.ppt.slide5_inventory import Slide5Inventory
from backend.models.domain import InventoryRecord
from backend.utils.exceptions import PPTGenerationError
from backend.ppt.placeholder_mapper import get_shape_alt_text

def test_slide5_inventory():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    t_shape = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(2), Inches(2))
    t_shape.name = "{{inventory_table}}"
    
    table = t_shape.table
    table.cell(0, 0).text = "Customer"
    table.cell(0, 1).text = "Service"
    table.cell(0, 2).text = "Links"
    
    s5 = Slide5Inventory(prs)
    s5.slide = slide # force to our slide since index might mismatch in raw prs
    
    records = [
        InventoryRecord(customer_name="ASG India", service_type="ILL", link_count=5),
        InventoryRecord(customer_name="ASG India", service_type="MPLS", link_count=3)
    ]
    
    s5.populate(records)
    
    assert len(table.rows) == 3
    assert table.cell(1, 1).text == "ILL"
    assert table.cell(2, 1).text == "MPLS"

def test_slide5_missing_placeholder():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    s5 = Slide5Inventory(prs)
    s5.slide = slide
    
    with pytest.raises(PPTGenerationError):
        s5.populate([])
