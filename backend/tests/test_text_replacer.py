from pptx import Presentation
from pptx.util import Inches
from backend.ppt.text_replacer import replace_text_in_slide

def test_text_replacer():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    txBox.text_frame.text = "Hello {{customer_name}}"
    
    replace_text_in_slide(slide, {"{{customer_name}}": "ASG India"})
    
    assert "ASG India" in txBox.text_frame.text
    assert "{{customer_name}}" not in txBox.text_frame.text
