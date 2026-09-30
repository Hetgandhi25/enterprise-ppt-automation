import pytest
from pptx import Presentation
from pptx.util import Inches
from backend.ppt.slide7_location_sla import Slide7LocationSLA
from backend.utils.exceptions import PPTGenerationError
import tempfile
from pathlib import Path
from PIL import Image

def create_dummy_image(path):
    img = Image.new('RGB', (10, 10), color = 'red')
    img.save(path)

def test_slide7_location_sla():
    prs = Presentation()
    for _ in range(7):
        prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides[6]
    
    with tempfile.TemporaryDirectory() as tmpdir:
        dummy_pic = Path(tmpdir) / "dummy.png"
        new_pic = Path(tmpdir) / "new.png"
        create_dummy_image(dummy_pic)
        create_dummy_image(new_pic)
        
        # Add a dummy shape so that the picture becomes index 1
        slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(1), Inches(1))
        # Add a picture to the slide
        slide.shapes.add_picture(str(dummy_pic), Inches(1), Inches(1))
        
        s7 = Slide7LocationSLA(prs)
        s7.slide = slide
        
        s7.populate(new_pic)
        
        # Verify the picture was replaced (it should still have 1 picture shape)
        pic_count = sum(1 for shape in slide.shapes if getattr(shape, 'shape_type', None) == 13)
        assert pic_count == 1

def test_slide7_missing_placeholder():
    prs = Presentation()
    for _ in range(7):
        prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides[6]
    s7 = Slide7LocationSLA(prs)
    s7.slide = slide
    
    with tempfile.TemporaryDirectory() as tmpdir:
        new_pic = Path(tmpdir) / "new.png"
        create_dummy_image(new_pic)
        
        # Should raise an error because there are no pictures to replace
        with pytest.raises(PPTGenerationError):
            s7.populate(new_pic)
