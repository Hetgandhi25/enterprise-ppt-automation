import pytest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from backend.ppt.slide6_performance import Slide6Performance
from backend.utils.exceptions import PPTGenerationError

def test_slide6_performance(tmp_path):
    prs = Presentation()
    for _ in range(4):
        prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides[1]
    
    from PIL import Image
    import tempfile
    
    with tempfile.TemporaryDirectory() as tmpdir:
        img_path = Path(tmpdir) / "dummy.png"
        img = Image.new('RGB', (10, 10), color = 'red')
        img.save(img_path)
        
        slide.shapes.add_picture(str(img_path), Inches(1), Inches(1))
        slide.shapes.add_picture(str(img_path), Inches(3), Inches(1))
    
        s6 = Slide6Performance(prs)
        s6.slide = slide
        
        s6.populate(img_path, img_path)
        
        # original shapes replaced, should still have 2 pictures
        assert len(slide.shapes) == 2
        assert slide.shapes[0].shape_type == 13
        assert slide.shapes[1].shape_type == 13

def test_slide6_missing_placeholder(tmp_path):
    prs = Presentation()
    for _ in range(4):
        prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides[1]
    s6 = Slide6Performance(prs)
    s6.slide = slide
    
    with pytest.raises(PPTGenerationError):
        s6.populate(tmp_path / "dummy.png", tmp_path / "dummy.png")
