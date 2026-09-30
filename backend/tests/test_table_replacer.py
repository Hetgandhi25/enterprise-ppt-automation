import pytest
from pptx import Presentation
from pptx.util import Inches
from backend.ppt.table_replacer import replace_table_data
from backend.utils.exceptions import PPTGenerationError

def test_table_replacer():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    t_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(2), Inches(2))
    table = t_shape.table
    table.cell(0, 0).text = "Col1"
    table.cell(0, 1).text = "Col2"
    table.cell(1, 0).text = "{{c1}}"
    table.cell(1, 1).text = "{{c2}}"
    
    data = [
        {"Col1": "Val1", "Col2": "Val2"},
        {"Col1": "Val3", "Col2": "Val4"}
    ]
    
    replace_table_data(t_shape, data, ["Col1", "Col2"])
    
    assert len(table.rows) == 3
    assert table.cell(1, 0).text == "Val1"
    assert table.cell(1, 1).text == "Val2"
    assert table.cell(2, 0).text == "Val3"
    assert table.cell(2, 1).text == "Val4"

def test_table_replacer_not_a_table():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(1), Inches(1))
    
    with pytest.raises(PPTGenerationError):
        replace_table_data(shape, [], [])
