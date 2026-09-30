from pptx import Presentation
from pptx.util import Inches
from backend.ppt.placeholder_mapper import find_placeholder, validate_placeholder, get_shape_alt_text

def test_placeholder_mapper():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(1), Inches(1))
    shape.name = "{{test_placeholder}}"
    
    assert get_shape_alt_text(shape) == "{{test_placeholder}}"
    
    found = find_placeholder(slide, "{{test_placeholder}}")
    assert found is not None
    assert found.name == "{{test_placeholder}}"
    
    assert validate_placeholder(slide, "{{test_placeholder}}") is True
    assert validate_placeholder(slide, "{{non_existent}}") is False
