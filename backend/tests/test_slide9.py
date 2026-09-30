import pytest
from pptx import Presentation
from pptx.util import Inches
from backend.ppt.slide9_incident import Slide9Incident
from backend.models.domain import IncidentRecord
from backend.utils.exceptions import PPTGenerationError

def test_slide9_incident():
    prs = Presentation()
    for _ in range(4):
        prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides[3]
    
    t_shape = slide.shapes.add_table(2, 4, Inches(1), Inches(1), Inches(2), Inches(2))
    t_shape.name = "{{incident_summary_table}}"
    table = t_shape.table
    table.cell(0, 0).text = "Location"
    table.cell(0, 1).text = "Cable issue"
    table.cell(0, 2).text = "Backhaul Impacted"
    table.cell(0, 3).text = "Electricity issue"
    
    s9 = Slide9Incident(prs)
    s9.slide = slide
    
    records = [IncidentRecord(customer_name="ASG India", location="Mumbai", cable_issue=2, backhaul_impacted=1, electricity_issue=0)]
    s9.populate(records)
    
    assert len(table.rows) == 2
    assert table.cell(1, 0).text == "Mumbai"
    assert table.cell(1, 1).text == "2"

def test_slide9_missing_placeholder():
    prs = Presentation()
    for _ in range(4):
        prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides[3]
    s9 = Slide9Incident(prs)
    s9.slide = slide
    
    # Should run cleanly without raising exception since it gracefully ignores missing table
    s9.populate([])
