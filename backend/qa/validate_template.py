import sys
import yaml
from pathlib import Path
from pptx import Presentation
from backend.ppt.shape_resolver import find_shape_by_id

def run_validation():
    manifest_path = Path("backend/templates/template_manifest.yaml")
    template_path = Path("backend/templates/MBR_Template_v1.pptx")
    
    with open(manifest_path, 'r', encoding='utf-8') as f:
        manifest = yaml.safe_load(f)
        
    prs = Presentation(template_path)
    slides_manifest = manifest.get("slides", {})
    
    passed = True
    errors = []
    
    for slide_key, mappings in slides_manifest.items():
        try:
            slide_idx = int(slide_key.split('_')[1]) - 1
        except:
            continue
            
        if slide_idx >= len(prs.slides):
            errors.append(f"Slide index {slide_idx} out of bounds.")
            passed = False
            continue
            
        slide = prs.slides[slide_idx]
        
        if mappings.get("pass"):
            continue
            
        for logical_name, shape_id_str in mappings.items():
            if logical_name == "pass": continue
            
            clean_id = str(shape_id_str).replace('text_', '').replace('table_', '').replace('chart_', '')
            shape = find_shape_by_id(slide, clean_id)
            
            if not shape:
                errors.append(f"{slide_key}: Mapping '{logical_name}' -> Shape ID {clean_id} NOT FOUND.")
                passed = False
                continue
                
            # Verify type
            if 'text' in shape_id_str and not shape.has_text_frame:
                errors.append(f"{slide_key}: {logical_name} -> Expected text frame, but shape has none.")
                passed = False
            elif 'table' in shape_id_str and not shape.has_table:
                errors.append(f"{slide_key}: {logical_name} -> Expected table, but shape has none.")
                passed = False
            elif 'chart' in shape_id_str and not shape.has_chart:
                errors.append(f"{slide_key}: {logical_name} -> Expected chart, but shape has none.")
                passed = False

    if passed:
        print("Template Manifest Validation: PASS")
    else:
        print("Template Manifest Validation: FAIL")
        for e in errors:
            print("  - " + e)
            
    return passed

if __name__ == "__main__":
    if not run_validation():
        sys.exit(1)
